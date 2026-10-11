from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"F:\game\Описание сайта Код доступа Эпохи.docx")
NAVY = "123B43"
PALE = "EEF5F4"
LIGHT = "F5F7F7"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{edge}")
        node = tc_mar.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_numbered(doc, lead, body):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(lead)
    r.bold = True
    p.add_run(body)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(10)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in (("Heading 1", 18, 14, 6), ("Heading 2", 14, 10, 4), ("Heading 3", 12, 8, 3)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[style_name].font.name = "Aptos"
    styles[style_name].font.size = Pt(11)
    styles[style_name].paragraph_format.space_after = Pt(3)

# Title and introduction
p = doc.add_paragraph(style="Title")
p.add_run("Описание образовательного сайта Код доступа Эпохи")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
r = subtitle.add_run("Игровая логика система оценивания и проверка качества")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(55, 67, 70)

p = doc.add_paragraph()
p.add_run("Назначение документа. ").bold = True
p.add_run(
    "Документ описывает устройство образовательной веб игры, последовательность прохождения, "
    "правила начисления баллов, таймеры, достижения, сохранение прогресса и автоматические тесты. "
    "Сайт рассчитан на командную работу и знакомит игроков с историей чисел и систем счисления."
)

doc.add_heading("Общая концепция", level=1)
doc.add_paragraph(
    "Код доступа Эпохи представляет собой интерактивную экспедицию из пяти этапов. Команда изучает "
    "древние способы записи чисел, выполняет переводы между системами счисления, исследует делители, "
    "двоичные разряды и ограничения восьмибитного диапазона. За прохождение этапов команда получает "
    "баллы, артефакты и достижения. Максимальный результат составляет 100 баллов."
)

doc.add_heading("Начало игры", level=2)
for item in (
    "Команда вводит название длиной до 24 символов.",
    "Можно указать одного или двух участников; их имена используются только в итоговой карточке.",
    "Перед запуском показывается инструктаж с правилами и фактами об истории чисел.",
    "Общий таймер запускается только после подтверждения начала экспедиции.",
    "Ответы, баллы и открытые этапы автоматически сохраняются в браузере.",
):
    add_bullet(doc, item)

doc.add_heading("Этапы экспедиции", level=1)
stages = [
    ("Взлом протокола", "Найти ошибку в историческом шифре, определить нарушенное правило и сформулировать гипотезу.", "Свиток исследователя"),
    ("Архив НИИ Времени", "Изучить системы счисления Египта, Вавилона, Рима и Древней Руси.", "Папирус числа"),
    ("Компиляция", "Записать 753 римскими и двоичными символами, исследовать делители 60 и 100.", "Табличка переводов"),
    ("Сборка ядра", "Определить двоичную ёмкость числа 3600 и проверить диапазон восьми бит.", "Печать совместимости"),
    ("Финальный блиц", "Выполнить два итоговых задания и ответить на пять вопросов за 90 секунд.", "Ключ временного ядра"),
]
for index, (name, body, artifact) in enumerate(stages, 1):
    p = add_numbered(doc, f"{name}. ", body)
    p.add_run(f" Артефакт: {artifact}.").italic = True

doc.add_heading("Навигация и разделы", level=1)
sections = [
    ("Главная", "текущее задание, шкала этапов и переходы между открытыми этапами"),
    ("Карта эпох", "цели, статус прохождения и артефакт каждого этапа"),
    ("Паспорт числа", "изученные системы, переводы, делители и результаты проверок"),
    ("Результаты", "баллы, попытки, подсказки, достижения, участники и итоговый статус"),
]
for name, body in sections:
    p = doc.add_paragraph()
    p.add_run(f"{name}. ").bold = True
    p.add_run(body.capitalize() + ".")
doc.add_paragraph(
    "Закрытые этапы недоступны до выполнения обязательной части предыдущего этапа. При переходе между "
    "разделами текущие ответы и прогресс сохраняются без перезагрузки страницы."
)

doc.add_heading("Система оценивания", level=1)
doc.add_paragraph(
    "Итог формируется из обязательных заданий, дополнительных задач и финального блица. "
    "Общий счёт ограничен 100 баллами."
)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(3.9)
table.columns[1].width = Inches(1.35)
table.columns[2].width = Inches(1.55)
headers = ("Источник баллов", "Количество", "Максимум")
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, NAVY)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_margins(cell)
rows = [
    ("Обязательные задания", "13", "70 баллов"),
    ("Засчитываемые дополнительные задания", "5 из 7", "25 баллов"),
    ("Финальный блиц", "5 вопросов", "5 баллов"),
    ("Итог", "", "100 баллов"),
]
for row_index, values in enumerate(rows, 1):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value
        cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cells[i])
        if row_index % 2 == 0:
            set_cell_shading(cells[i], PALE)
        if row_index == len(rows):
            for run in cells[i].paragraphs[0].runs:
                run.bold = True
set_table_borders(table)

doc.add_heading("Попытки и подсказки", level=2)
doc.add_paragraph(
    "Обычное задание обычно оценивается в 5 баллов. Первая правильная попытка сохраняет полную стоимость, "
    "вторая уменьшает награду на 1 балл, а третья и последующие - на 2 балла. Минимальная награда за "
    "выполненное обычное задание составляет 1 балл."
)
doc.add_paragraph(
    "Практическое задание четвёртого этапа стоит 10 баллов. Вторая попытка уменьшает результат на 2 балла, "
    "третья и последующие - на 4 балла. Минимальная награда составляет 2 балла."
)
doc.add_paragraph(
    "Открытие подсказки дополнительно уменьшает результат на 1 балл в обычном задании и на 2 балла в "
    "практическом. Штрафы за попытки и подсказку складываются."
)

doc.add_heading("Условия успешного завершения", level=2)
for item in (
    "выполнены все обязательные задания",
    "пройдены все пять этапов",
    "собраны пять артефактов",
    "набрано не менее 60 баллов",
):
    add_bullet(doc, item)
doc.add_paragraph(
    "Команда, которая собрала все артефакты и набрала не менее 90 баллов, получает достижение Полный допуск."
)

doc.add_heading("Режим восстановления", level=2)
doc.add_paragraph(
    "Если баллы были потеряны из-за нескольких попыток или подсказок, на странице результатов доступно "
    "повторное выполнение подходящих заданий. За обычное задание можно вернуть до 2 баллов, за практическое "
    "- до 4 баллов. Каждое задание восстанавливается только один раз, а итог не может превысить 100 баллов."
)

doc.add_heading("Достижения", level=1)
achievements = [
    ("Архивист эпох", "выполнены задания всех четырёх цивилизаций"),
    ("Безошибочный дешифровщик", "основные задания этапов 2 и 3 выполнены с первой попытки без подсказок"),
    ("Мастер шестидесяти", "исследованы делители 60 и выполнено дополнительное задание Вавилона"),
    ("Хранитель разрядов", "проверены двоичная запись 753 и диапазон восьми бит"),
    ("Полный допуск", "собраны все артефакты и набрано не менее 90 баллов"),
]
for name, condition in achievements:
    p = doc.add_paragraph()
    p.add_run(f"{name}. ").bold = True
    p.add_run(condition.capitalize() + ".")

doc.add_heading("Проверка ответов", level=1)
doc.add_paragraph(
    "Универсальный движок поддерживает целые числа, текст, одиночный и множественный выбор, римские числа, "
    "двоичные записи и списки чисел. Перед сравнением ответ нормализуется, чтобы не считать ошибкой "
    "различия в регистре, пробелах и допустимом оформлении."
)
normalization = [
    ("+007", "7"),
    ("dcc liii", "DCCLIII"),
    ("0b1011 1100_01₂", "1011110001"),
    ("6, 2; 3 2", "2,3,6"),
    ("НОЛЬ и Ёлка", "ноль и елка"),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(3.1)
table.columns[1].width = Inches(3.7)
for i, text in enumerate(("Исходный ввод", "Нормализованное значение")):
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, NAVY)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
    set_cell_margins(cell)
for row_index, values in enumerate(normalization, 1):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value
        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cells[i])
        if row_index % 2 == 0:
            set_cell_shading(cells[i], PALE)
set_table_borders(table)

doc.add_heading("Таймеры и сохранение прогресса", level=1)
add_bullet(doc, "Общий таймер экспедиции составляет 90 минут.")
add_bullet(doc, "Финальный блиц длится 90 секунд.")
add_bullet(doc, "На мобильных устройствах текущий таймер дополнительно отображается в хедере.")
add_bullet(doc, "После истечения общего времени формы блокируются, но накопленный результат сохраняется.")
add_bullet(doc, "Состояние игры хранится локально в браузере и восстанавливается после перезагрузки.")
doc.add_paragraph(
    "Полный сброс очищает название команды, участников, ответы, попытки, подсказки, баллы, таймеры, "
    "артефакты и достижения, после чего открывается стартовый экран."
)

doc.add_heading("Автоматическое тестирование", level=1)
doc.add_paragraph(
    "В проекте настроены модульные и smoke тесты на базе встроенного тестового раннера Node.js. "
    "Текущий набор содержит семь проверок."
)
tests = [
    "банк содержит 13 обязательных и 7 дополнительных вопросов",
    "обязательные задания охватывают все пять этапов",
    "существуют экраны пяти этапов и сценарий публикации GitHub Pages",
    "сохраняется не более двух имён участников",
    "баллы начисляются только за пять из семи дополнительных заданий",
    "полный сброс удаляет команду, участников и игровой прогресс",
    "нормализация и валидаторы корректно обрабатывают поддерживаемые форматы ответов",
]
for item in tests:
    add_bullet(doc, item.capitalize() + ".")
doc.add_paragraph(
    "Последний полный запуск завершился успешно: пройдено 7 из 7 тестов, ошибок и пропущенных проверок нет. "
    "Production сборка Webpack также формируется успешно."
)

doc.add_heading("Адаптивность и доступность", level=1)
doc.add_paragraph(
    "Интерфейс адаптирован для мобильных устройств, планшетов, ноутбуков и широких экранов. Колонки, "
    "карточки, формы, навигация, таблицы и модальные окна перестраиваются под доступную ширину. "
    "Интерактивные элементы сохраняют удобный размер нажатия, изображения не растягивают контейнеры, "
    "а длинные названия и ответы переносятся без горизонтальной прокрутки."
)
doc.add_paragraph(
    "Для клавиатурной навигации предусмотрен видимый фокус на кнопках, ссылках и полях. Заголовки экранов "
    "получают программный фокус при переходе, что помогает пользователям вспомогательных технологий."
)

doc.add_heading("Итог", level=1)
doc.add_paragraph(
    "Сайт объединяет учебные задания, игровую прогрессию и автоматическую проверку ответов. Полный маршрут "
    "позволяет набрать 100 баллов, фиксирует качество прохождения через попытки и подсказки, выдаёт пять "
    "артефактов и пять достижений, а также сохраняет состояние команды между сеансами."
)

# Footer with document name and a dynamic page field.
footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(6)
run = footer_p.add_run("Код доступа Эпохи   •   ")
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(90, 90, 90)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer_p._p.append(fld)

doc.core_properties.title = "Описание образовательного сайта Код доступа Эпохи"
doc.core_properties.subject = "Игровая логика система оценивания и тестирование"
doc.core_properties.author = ""
doc.core_properties.keywords = "образовательная игра, системы счисления, оценивание, тестирование"

doc.save(OUT)
print(OUT)
